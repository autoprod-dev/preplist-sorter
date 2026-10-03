# Headless Chromium tests for Prepline v3 (Gerard's 2026-10-02 baseline + Prepline/Pro layer + tweaks).
# Usage: python3 test_v3.py URL [shots_dir] [out_dir] [--expect-version X]
import sys, os, json, re, time
from urllib.parse import urlparse
from playwright.sync_api import sync_playwright
import openpyxl
URL = sys.argv[1]
SHOTS = sys.argv[2] if len(sys.argv) > 2 and not sys.argv[2].startswith('--') else '/workspace/prepline-v2-shots'
OUT = sys.argv[3] if len(sys.argv) > 3 and not sys.argv[3].startswith('--') else '/workspace/prepline-v3-test-out'
EXPECT_VER = sys.argv[sys.argv.index('--expect-version')+1] if '--expect-version' in sys.argv else None
OFFLINE = '--offline' in sys.argv   # file:// build, ALL network blocked; any request attempt fails the run
LIVE = '--live' in sys.argv or OFFLINE   # public build: LICENSE_CONFIG mode 'off' (no Pro UI, no Lemon Squeezy)
S = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'samples') + '/'
os.makedirs(SHOTS, exist_ok=True); os.makedirs(OUT, exist_ok=True)
results = []; ext = {}; net_attempts = []
def check(name, ok, info=''):
    results.append({'test': name, 'ok': bool(ok), 'info': str(info)[:600]}); print(('PASS ' if ok else 'FAIL ') + name + ('  | ' + str(info)[:300] if info else ''))
def track(ctx):
    def on_req(r):
        h = urlparse(r.url).hostname or ''
        if OFFLINE and not r.url.startswith(('file:', 'data:', 'blob:')): net_attempts.append(r.url[:140])
        if h not in ('localhost', '127.0.0.1') and not r.url.startswith(('data:', 'blob:', 'file:')):
            ext.setdefault(h, set()).add(r.method + ' ' + r.url.split('?')[0][:140])
    ctx.on('request', on_req)
def new_ctx(b, w=1280, h=900, **kw):
    ctx = b.new_context(viewport={'width': w, 'height': h}, accept_downloads=True, **kw); track(ctx)
    if OFFLINE:
        ctx.set_offline(True)
        ctx.route(re.compile(r'^(https?|wss?|ftp)://.*'), lambda route: (net_attempts.append('ROUTED ' + route.request.url[:140]), route.abort()))
    ls = []
    ctx.route(re.compile(r'.*lemonsqueezy\.com/checkout.*'), lambda route: (ls.append(route.request.url), route.abort()))
    return ctx, ls
def page_with_errs(ctx):
    pg = ctx.new_page(); errs = []
    pg.on('pageerror', lambda e: errs.append('pageerror: ' + str(e)))
    pg.on('console', lambda m: m.type == 'error' and errs.append('console: ' + m.text[:200]))
    return pg, errs
def up(pg, files):
    pg.evaluate("document.getElementById('fileInput').value=''; document.getElementById('statusMsg').textContent=''"); pg.set_input_files('#fileInput', files)
def wait_status(pg, pat, t=20000):
    pg.wait_for_function("p => new RegExp(p).test(document.getElementById('statusMsg').textContent)", arg=pat, timeout=t)
    return pg.inner_text('#statusMsg')
def capture_print(pg, fn_js):
    pg.evaluate("window.__prints=[]; window.printViaIframe = h => window.__prints.push(h);")
    pg.evaluate(fn_js); pg.wait_for_function("window.__prints.length>0", timeout=10000)
    return pg.evaluate("window.__prints[window.__prints.length-1]")
def export_view(pg, view, fname):
    pg.evaluate(f"setView('{view}')"); pg.wait_for_timeout(300)
    btn = pg.locator('button:has-text("📊 Export to Excel (editable)")')
    n = btn.count()
    with pg.expect_download() as dl: btn.first.click()
    path = os.path.join(OUT, fname); dl.value.save_as(path)
    return path, n, btn.first.inner_text()
def xlsx_info(path):
    wb = openpyxl.load_workbook(path)
    info = {}
    for ws in wb.worksheets:
        vals = [[c.value for c in row] for row in ws.iter_rows()]
        info[ws.title] = {'images': len(getattr(ws, '_images', [])), 'rows': ws.max_row,
                          'first': vals[0][0] if vals else None,
                          'formulas': sum(1 for row in vals for v in row if isinstance(v, str) and v.startswith('=')),
                          'header': next((r for r in vals if r and r[0] == 'Item'), None)}
    return info

with sync_playwright() as p:
    b = p.chromium.launch()
    # ---------- A. load ----------
    ctx, ls = new_ctx(b); pg, errs = page_with_errs(ctx)
    pg.goto(URL); pg.wait_for_load_state('networkidle')
    ver = pg.get_attribute('meta[name="prepline-version"]', 'content')
    if OFFLINE:
        pg.wait_for_function('window.pdfjsLib', timeout=15000)
        check('offline: pdf.js + worker loaded from the bundled copy', pg.evaluate('window.pdfjsWorkerMode') == 'worker' and pg.evaluate('typeof XLSX') == 'object', pg.evaluate('window.pdfjsWorkerMode'))
        check('offline: no web fonts referenced (system fonts)', not pg.evaluate("[...document.querySelectorAll('link[href]')].some(l=>/^https?:/.test(l.href))"))
    check('load: title is Prepline', pg.title().startswith('Prepline'), pg.title())
    check('load: version string present' + (f' = {EXPECT_VER}' if EXPECT_VER else ''), ver and (not EXPECT_VER or ver == EXPECT_VER) and ('v' + ver) in pg.inner_text('#appVersion'), ver)
    if not LIVE:
        check('load: TEST MODE badge visible', pg.is_visible('#plTestBadge') and 'TEST MODE' in pg.inner_text('#plTestBadge'), pg.inner_text('#plTestBadge'))
        check('load: Upgrade + Enter licence key buttons', pg.is_visible('#proUpgradeBtn') and pg.is_visible('#licenceKeyBtn'))
    else:
        check('live: no TEST MODE badge', pg.locator('#plTestBadge').count() == 0 or not pg.is_visible('#plTestBadge'))
        check('live: no Upgrade / Enter licence key buttons, no Pro panel', not pg.is_visible('#proUpgradeBtn') and not pg.is_visible('#licenceKeyBtn') and not pg.is_visible('#proPanel'))
        check('live: no noindex, title has no (TEST)', pg.locator('meta[name=robots]').count() == 0 and '(TEST)' not in pg.title(), pg.title())
    check('load: no JS errors', not errs, errs)
    pg.screenshot(path=f'{SHOTS}/00-load-desktop.png')
    # ---------- B. PDF sort + method capture ----------
    up(pg, S + 'sample-recipes.pdf')
    st = wait_status(pg, 'Sorted|Could not')
    rows = pg.evaluate('ROWS.length'); crops = pg.evaluate('Object.fromEntries(Object.entries(methodCrops).map(([k,v])=>[k,v.length]))')
    secs = pg.evaluate('[...new Set(ROWS.map(r=>r.section))]')
    check('PDF: prep list sorted (9 items, 2 dishes)', rows == 9 and 'Sorted' in st, st)
    check('PDF: items split into sections', len(secs) >= 2, secs)
    check('PDF: method crop captured for each dish', len(crops) == 2 and all(v >= 1 for v in crops.values()), crops)
    check('PDF: "View original" button per dish', pg.locator('button:has-text("View original")').count() == 2)
    pg.evaluate("setView('dish')"); pg.wait_for_timeout(300)
    pg.evaluate("document.querySelector('#resultsArea .viewtoggle').scrollIntoView()"); pg.wait_for_timeout(300)
    pg.screenshot(path=f'{SHOTS}/01-prep-list-by-dish-desktop.png', full_page=False)
    h0 = capture_print(pg, "printDish(encodeURIComponent(Object.keys(methodCrops)[0]))")
    check('v3.2.5: fresh storage, print toggle shows No and the dish print has no method captures', pg.evaluate("[includeMethodInPrint, localStorage.getItem('preplinePref_printIncludeMethod_v1'), document.querySelector('#printMethodToggle [data-pmt=no]').getAttribute('aria-pressed')]") == [False, None, 'true'] and 'Method reference' not in h0 and '<img' not in re.sub(r'<style.*?</style>', '', h0, flags=re.S), len(h0))
    pg.click('#printMethodToggle [data-pmt=yes]'); pg.wait_for_timeout(200)   # switch to Yes for the method print checks below
    html = capture_print(pg, "printDish(encodeURIComponent(Object.keys(methodCrops)[0]))")
    check('method (toggle Yes): printed dish sheet has the Method crop below the prep list', 'Method reference' in html and 'data:image/jpeg' in html and html.index('Method reference') > html.index('<table'), len(html))
    pp = ctx.new_page(); pp.set_viewport_size({'width': 900, 'height': 1200}); pp.set_content(html); pp.emulate_media(media='print'); pp.wait_for_timeout(800)
    pp.screenshot(path=f'{SHOTS}/02-method-captured-below-prep-list.png', full_page=True); pp.close()
    for fn, lbl in [('printAllSections()', 'sections'), ('printAllDays()', 'days')]:
        h = capture_print(pg, fn); check(f'print all {lbl}: builds with method crops', 'Method reference' in h, len(h))
    labels = pg.locator('button:has-text("Print all")').all_inner_texts()
    check('tweak 1: print button labels', labels == ['🖨️ Print all sections (one page complete)', '🖨️ Print all prep days (one page complete)'], labels)
    # ---------- C. reload ----------
    errs.clear(); pg.wait_for_timeout(1000); pg.reload(); pg.wait_for_load_state('networkidle'); pg.wait_for_timeout(800)
    txt = pg.inner_text('#resultsArea') if pg.locator('#resultsArea').count() else pg.inner_text('body')
    check('reload: list restored, no blank page, no errors', pg.evaluate('ROWS.length') == 9 and 'Pumpkin' in txt and not errs, errs)
    check('reload: method crops survive reload', len(pg.evaluate('Object.keys(methodCrops)')) == 2)
    # ---------- D. image-only inputs (does anything do OCR / call out?) ----------
    before = {h: set(v) for h, v in ext.items()}
    up(pg, S + 'recipe-screenshot-only.pdf'); st2 = wait_status(pg, 'Could not|Sorted', 30000)
    check('screenshot-only PDF (no text layer): not read, no OCR', 'Could not' in st2, st2)
    up(pg, S + 'recipe-screenshot.png'); pg.wait_for_timeout(2500)
    st3 = pg.inner_text('#statusMsg')
    check('PNG screenshot upload: not supported (accept=.xlsx,.xls,.pdf)', pg.get_attribute('#fileInput', 'accept') == '.xlsx,.xls,.pdf' and 'Sorted' not in st3, st3)
    newreq = {h: v - before.get(h, set()) for h, v in ext.items() if v - before.get(h, set())}
    check('image inputs: no new outside requests', not newreq, newreq)
    # ---------- E. licence ----------
    if not LIVE:
        pg.click('#licenceKeyBtn'); pg.fill('#plKey', 'NOT-A-KEY'); pg.click('#plActivate'); pg.wait_for_timeout(1500)
        check('licence: made-up key refused', 'pl-msg err' in (pg.get_attribute('#plLicMsg', 'class') or ''), pg.inner_text('#plLicMsg'))
        pg.fill('#plKey', 'TEST-PRO-YEARLY'); pg.click('#plActivate'); pg.wait_for_timeout(800)
        check('licence: TEST-PRO-YEARLY unlocks Pro', pg.evaluate('AutoprodLicence.isPro()') and 'Pro' in pg.inner_text('#plTestBadge'), pg.inner_text('#plTestBadge'))
        pg.keyboard.press('Escape'); pg.wait_for_timeout(300)
        check('licence: Pro panel shows (templates, stations, kitchens)', 'Saved templates' in pg.inner_text('#proPanel') and not pg.is_visible('#licenceKeyBtn'))
        pg.screenshot(path=f'{SHOTS}/03-pro-unlocked-desktop.png')
        pg.wait_for_timeout(1000); pg.reload(); pg.wait_for_load_state('networkidle')
        check('licence: Pro survives reload', pg.evaluate('AutoprodLicence.isPro()'))
        # append (Pro)
        up(pg, S + 'sample-recipes.pdf'); wait_status(pg, 'Sorted')
        pg.evaluate("document.getElementById('statusMsg').textContent=''"); pg.set_input_files('#appendInput', S + 'sample-recipes.pdf'); st4 = wait_status(pg, 'Appended|Could not', 30000)
        check('Pro: append another PDF keeps list + method crops', 'Appended' in st4 and pg.evaluate('ROWS.length') == 18 and len(pg.evaluate('Object.keys(methodCrops)')) == 4, st4[:160])
        for k in ['TEST-PRO-MONTHLY', 'TEST-PRO-LIFETIME']:
            r = pg.evaluate(f"AutoprodLicence.deactivate().then(()=>AutoprodLicence.activate('{k}')).then(r=>r.ok && AutoprodLicence.plan())")
            check(f'licence: {k} unlocks', r in ('monthly', 'lifetime'), r)
        pg.evaluate('AutoprodLicence.deactivate()'); pg.wait_for_timeout(300)
        check('licence: deactivate returns to free', not pg.evaluate('AutoprodLicence.isPro()'))
        # ---------- F. checkout links ----------
        want = {'monthly': 'https://autoprod-dev.lemonsqueezy.com/checkout/buy/26141028-ac72-45a5-95d5-07ebf96e9a2e?enabled=2179964',
                'yearly': 'https://autoprod-dev.lemonsqueezy.com/checkout/buy/bbdeab5d-7ff3-48e8-88d1-5d0fec3c620c?enabled=2179971',
                'lifetime': 'https://autoprod-dev.lemonsqueezy.com/checkout/buy/d69e7ed5-d47d-40af-8d17-a7e599a89314?checkout%5Bdiscount_code%5D=LAUNCH50'}
        pg.click('#proUpgradeBtn'); pg.wait_for_timeout(300)
        got = {}
        for plan in ['monthly', 'yearly', 'lifetime']:
            with ctx.expect_page() as pop: pg.click(f'#plUpgradeDialog [data-plan="{plan}"]')
            pop.value.wait_for_timeout(500); got[plan] = ls[-1] if ls else None; pop.value.close()
        for plan in want: check(f'checkout: Buy {plan} opens the test-mode LS checkout', got[plan] == want[plan], got[plan])
        check('checkout: page stays on the app', pg.url.startswith(URL.split('#')[0].rsplit('/', 1)[0]))
        pg.keyboard.press('Escape')
    else:
        # live build: Pro switched off (mode 'off'): no licence, no checkout, no Pro tools
        r = pg.evaluate("AutoprodLicence.activate('TEST-PRO-YEARLY').then(r=>[r.ok, AutoprodLicence.isPro(), AutoprodLicence.isTestMode()])")
        check('live: Pro is off — a TEST key does not unlock anything, not test mode', r == [False, False, False], r)
        pg.evaluate("AutoprodLicence.openUpgrade(); AutoprodLicence.openLicenceDialog(); AutoprodLicence.buy('yearly')"); pg.wait_for_timeout(500)
        check('live: Upgrade / licence / Buy calls do nothing (no dialog, no checkout tab)', pg.locator('dialog[open]').count() == 0 and len(ctx.pages) == 1 and not ls, (pg.locator('dialog[open]').count(), len(ctx.pages), ls))
        vis = pg.evaluate("['proUpgradeBtn','licenceKeyBtn','proPanel','appendRow','appendPdfBtn','plTestBadge'].map(id=>{const e=document.getElementById(id); return [id, !!(e && e.offsetParent)];})")
        check('live: with a list loaded, no Upgrade, licence key, Pro panel, Append PDF or TEST badge is visible', not any(v for _, v in vis), vis)
        body = pg.inner_text('body')
        check('live: page text has no TEST MODE, Upgrade, Buy, Lemon Squeezy or test card wording', not re.search(r'TEST MODE|Upgrade|Buy (monthly|yearly|lifetime)|Lemon Squeezy|4242|licence key', body, re.I), re.findall(r'.{30}(?:TEST MODE|Upgrade|Lemon Squeezy|4242|licence key).{30}', body, re.I)[:3])
    # ---------- G. exports: tab names + Source tab ----------
    def src_tabs_xml(path):
        import zipfile, html as _h
        x = zipfile.ZipFile(path).read('xl/workbook.xml').decode('utf-8')
        return [_h.unescape(n) for n in re.findall(r'<sheet\b[^>]*?\bname="([^"]*)"', x)]
    def excel_rule(names):
        out, used = [], set()
        for o in names:
            t = re.sub(r'[\\/?*\[\]:]', '-', o).strip("'")[:31] or 'Sheet'
            if t.lower() == 'history': t = 'History-1'
            f, n = t, 2
            while f.lower() in used: f = t[:31-len(f'-{n}')] + f'-{n}'; n += 1
            used.add(f.lower()); out.append(f)
        return out
    def rows_of(path):
        wb = openpyxl.load_workbook(path); out = []
        for ws in wb.worksheets:
            hdr = None
            for row in ws.iter_rows(values_only=True):
                if row and row[0] == 'Item': hdr = list(row); continue
                if hdr and row and row[0] and len(row) > 7 and row[7] is not None:
                    out.append((ws.title, row[0], row[7]))
        return out, wb
    def check_exports(src_label, expect_tabs, task_src):
        for view, fn in [('dish', f'{src_label}-by-dish.xlsx'), ('section', f'{src_label}-by-section.xlsx'), ('priority', f'{src_label}-by-day.xlsx')]:
            path, n, label = export_view(pg, view, fn); info = xlsx_info(path); exp[fn] = info
            st = pg.inner_text('#statusMsg')
            check(f'{src_label} {view}: own "📊 Export to Excel (editable)" button', n == 1 and label.strip() == '📊 Export to Excel (editable)', label)
            check(f'{src_label} {view}: editable cells, no images, Item + Source tab header', all(v['images'] == 0 and v['header'] and v['header'][7] == 'Source tab' for v in info.values()), list(info))
            rows, wb = rows_of(path)
            bad = [(t, it, sv) for (t, it, sv) in rows if task_src.get(it) != sv]
            check(f'{src_label} {view}: Source tab column = exact original name on every row ({len(rows)} rows)', rows and not bad, bad[:4])
            if view == 'dish':
                check(f'{src_label} by dish: tab names = original names (Excel-forced changes only), same order', src_tabs_xml(path) == expect_tabs, (src_tabs_xml(path), expect_tabs))
                exp_changes = sum(1 for a, b2 in zip(list(dict.fromkeys(task_src.values())), expect_tabs) if a != b2)
                check(f'{src_label} by dish: every forced rename is reported', (exp_changes == 0 and 'changed because Excel requires it' not in st) or (f'{exp_changes} tab name' in st), st[:300])
            else:
                check(f'{src_label} {view}: tabs are section/day names', all(t in [x[:31] for x in pg.evaluate("[...SECTIONS, ...WEEKDAYS, 'Unscheduled']")] for t in src_tabs_xml(path)), src_tabs_xml(path))
        return exp
    exp = {}
    # G1. plain multi-sheet workbook
    up(pg, S + 'sample-multisheet.xlsx'); stx = wait_status(pg, 'Sorted|Could not')
    check('xlsx: multi-sheet workbook sorted', 'Sorted' in stx, stx)
    src = src_tabs_xml(S + 'sample-multisheet.xlsx')
    tsrc = {r[0]: r[1] for r in pg.evaluate('ROWS.map(r=>[r.task, r.srcSheet])')}
    check_exports('multisheet', src, tsrc)
    check('export: Total needed is a live formula (editable)', sum(v['formulas'] for v in exp['multisheet-by-dish.xlsx'].values()) > 0)
    staff = [r[0].value for r in openpyxl.load_workbook(os.path.join(OUT, 'multisheet-by-dish.xlsx'))['Staff Meal'].iter_rows() if r[0].value in ('Chicken Curry', 'Steamed Rice')]
    check('a source tab with two dishes exports as one tab with both', staff == ['Chicken Curry', 'Steamed Rice'], staff)
    # G2. awkward tab names
    up(pg, S + 'awkward-tabs.xlsx'); stx = wait_status(pg, 'Sorted|Could not')
    awk = src_tabs_xml(S + 'awkward-tabs.xlsx')
    got_src = pg.evaluate('[...new Set(ROWS.map(r=>r.srcSheet))]')
    check('awkward xlsx: every original tab name stored exactly on its rows', sorted(got_src) == sorted(awk), got_src)
    tsrc = {r[0]: r[1] for r in pg.evaluate('ROWS.map(r=>[r.task, r.srcSheet])')}
    expect = excel_rule(awk)
    check('awkward: expected Excel-forced changes are only MAINS (case dup) and the 40-char name', [(a, b2) for a, b2 in zip(awk, expect) if a != b2] == [('MAINS', 'MAINS-2'), (awk[3], awk[3][:31])], list(zip(awk, expect)))
    check_exports('awkward', expect, tsrc)
    # round trip: Source tab column restores the exact originals, even renamed ones
    up(pg, os.path.join(OUT, 'awkward-by-section.xlsx')); wait_status(pg, 'Sorted|Could not')
    rt = pg.evaluate('[...new Set(ROWS.map(r=>r.srcSheet))]')
    check('round trip: re-uploading a By section export restores exact original tab names', sorted(rt) == sorted(awk), rt)
    path, _, _ = export_view(pg, 'dish', 'awkward-roundtrip-by-dish.xlsx')
    check('round trip: By dish export after re-upload gives the same tabs', sorted(src_tabs_xml(path)) == sorted(expect), src_tabs_xml(path))
    up(pg, os.path.join(OUT, 'awkward-by-dish.xlsx')); wait_status(pg, 'Sorted|Could not')
    rt2 = {r[0]: r[1] for r in pg.evaluate('ROWS.map(r=>[r.task, r.srcSheet])')}
    check('round trip: re-uploading the By dish export keeps every row\'s original tab', rt2 == tsrc, {k: v for k, v in rt2.items() if tsrc.get(k) != v})
    # sanitiser unit checks
    names = pg.evaluate("""(()=>{const u=new Set(), c=[]; const r=['Mains: Beef [Sat]','mains- beef -sat-',"'Quoted'",'History','A'.repeat(40),'A'.repeat(40),'Tab?/*\\\\x','  keep  spaces ',"Chef's"].map(n=>safeSheetName(n,u,c)); return {r, c:c.length};})()""")
    nm = names['r']
    ok = all(len(n) <= 31 and not re.search(r'[\[\]:*?/\\]', n) for n in nm) and len({n.lower() for n in nm}) == len(nm) and nm[2] == 'Quoted' and nm[3] == 'History-1' and nm[7] == '  keep  spaces ' and nm[8] == "Chef's" and names['c'] == 7
    check('sanitiser: changes only what Excel forces (spaces/apostrophes inside kept), reports each change', ok, names)
    # G3. text PDF: recipe titles as the tab names
    up(pg, S + 'fixes-and-titles.pdf'); stp = wait_status(pg, 'Sorted|Could not', 30000)
    titles = ['Braised Lamb Shoulder: 18 Hour [GF]', "Market Fish's Salsa Verde"]
    check('PDF: source name = recipe title exactly as printed', pg.evaluate('[...new Set(ROWS.map(r=>r.srcSheet))]') == titles, pg.evaluate('[...new Set(ROWS.map(r=>r.srcSheet))]'))
    tsrc = {r[0]: r[1] for r in pg.evaluate('ROWS.map(r=>[r.task, r.srcSheet])')}
    check_exports('pdf', excel_rule(titles), tsrc)
    # ---------- G4. the four fixes ----------
    rws = {r['task']: r for r in pg.evaluate('ROWS')}
    check('fix 1: merged "2.5 Kilograms Chicken thigh" split into qty + name', 'Chicken thigh' in rws and rws['Chicken thigh']['pdfFixedTotal'] == 2.5 and rws['Chicken thigh']['unit'] == 'kg', rws.get('Chicken thigh'))
    check('fix 1: merged "Lemon zest 200 Grams" (name first) split too', 'Lemon zest' in rws and rws['Lemon zest']['pdfFixedTotal'] == 200, rws.get('Lemon zest'))
    check('fix 2: butternut pumpkin (PDF) is Vegetables, not a sauce', rws['Butternut pumpkin']['section'] == 'Vegetables', rws['Butternut pumpkin']['section'])
    pg.evaluate("setView('dish')"); pg.wait_for_timeout(300)
    def row_text(task): return pg.locator(f'#resultsArea tr:has(td:has-text("{task}"))').first.inner_text()
    rw, jus = row_text('Red wine'), row_text('Lamb jus')
    check('fix 3: 1500 Milliliters shows as 1.5 L (never kg)', '1.5 L' in rw and 'kg' not in rw, rw)
    check('fix 3: 2 Liters shows as 2 L (never kg)', '2 L' in jus and 'kg' not in jus, jus)
    check('fix 4: status message says COOKED/finished, not "raw ingredient weights"', 'COOKED' in stp and 'Raw ingredient weights only' not in stp, stp[:200])
    ph = capture_print(pg, "printDish(encodeURIComponent(ROWS[0].dish))")
    check('fix 4: printed banner agrees (COOKED WEIGHT)', 'COOKED WEIGHT' in ph)
    hint_help = pg.evaluate("document.getElementById('helpTab').innerText + ' ' + document.body.innerHTML")
    check('fix 4: help text no longer says PDFs are read for raw ingredients', 'read for its raw ingredients' not in hint_help and 'finished/cooked weights' in hint_help)
    up(pg, S + 'awkward-tabs.xlsx'); wait_status(pg, 'Sorted')
    rws = {r['task']: r for r in pg.evaluate('ROWS')}
    check('fix 2: "Butternut pumpkin roast" (xlsx) is not a sauce', rws['Butternut pumpkin roast']['section'] != 'Sauce, Purees and Gels', rws['Butternut pumpkin roast']['section'])
    pg.evaluate("setView('dish')"); pg.wait_for_timeout(300)
    t1, t2, t3 = row_text('Red wine jus'), row_text('Butternut pumpkin roast'), row_text('Beef fillet portion')
    check('fix 3: xlsx 0.05 L x 80 = 4 L; 0.015 kg x 40 = 600 g; 0.18 kg x 80 = 14.4 kg', '4 L' in t1 and '600 g' in t2 and '14.4 kg' in t3, (t1, t2, t3))
    # ---------- G4b. v3.2.1 classifier ----------
    BJ, VEG, SAU, GAR, PRO = 'Big Jobs', 'Vegetables', 'Sauce, Purees and Gels', 'Garnishes, Herbs, Powders, Oils Etc.', 'Proteins'
    table = {'Meringue bake': BJ, 'Jasmine rice wash': BJ, 'Steamed rice': BJ, 'Risotto base': BJ, 'Couscous steep': BJ,
             'Fresh pasta dough': BJ, 'Gnocchi roll': BJ, 'Focaccia bake': BJ, 'Pavlova base bake': BJ, 'Lemon tart shell': BJ,
             'Chocolate mousse': BJ, 'Vanilla ice cream': BJ, 'Gelato base': BJ, 'Pancake batter': BJ, 'Cheesecake base': BJ,
             'Chicken liver parfait': PRO, 'Salmon mousse': PRO, 'Beef tartare': PRO, 'Puff pastry for beef wellington': PRO,
             'Cauliflower florets': VEG, 'Brussels sprouts halve': VEG, 'Trumpet mushroom slice': VEG, 'Green beans trim': VEG,
             'Breadcrumb for crumbing': GAR, 'Edible flowers pick': GAR, 'Rice wine vinegar': SAU, 'Red wine jus': SAU,
             'Raspberry gel': SAU, 'Unsalted butter': SAU, 'Butternut pumpkin': VEG, 'Beef rump portion': PRO, 'Potato mash': BJ}
    got = pg.evaluate("t => Object.fromEntries(Object.keys(t).map(k=>[k, categorize(k,'')]))", table)
    wrong = {k: (got[k], v) for k, v in table.items() if got[k] != v}
    check(f'classifier: {len(table)} items land in the expected section (desserts/starch → Big Jobs, no substring mis-hits)', not wrong, wrong)
    days = pg.evaluate("[classifyPrepStage('Swiss chard wash','Vegetables',null), classifyPrepStage('Pavlova base portion','Big Jobs',null)]")
    check('classifier: "Swiss chard wash" is a day-1 wash (not "char" cook); pavlova not treated as a pave', days == [1, 3], days)
    up(pg, S + 'sample-multisheet.xlsx'); wait_status(pg, 'Sorted')
    rr = {r['task']: r for r in pg.evaluate('ROWS')}
    check('classifier e2e: "Meringue bake" → Big Jobs, not flagged', rr['Meringue bake']['section'] == BJ and not rr['Meringue bake']['flagged'], rr['Meringue bake'])
    check('classifier e2e: "Jasmine rice wash" → Big Jobs, not flagged', rr['Jasmine rice wash']['section'] == BJ and not rr['Jasmine rice wash']['flagged'], rr['Jasmine rice wash'])
    up(pg, S + 'awkward-tabs.xlsx'); wait_status(pg, 'Sorted')
    # ---------- G5. source names survive reload, templates, append ----------
    before_src = pg.evaluate('ROWS.map(r=>r.srcSheet)')
    pg.wait_for_timeout(1000); pg.reload(); pg.wait_for_load_state('networkidle'); pg.wait_for_timeout(600)
    after_src = pg.evaluate('ROWS.map(r=>r.srcSheet)'); check('reload: every row keeps its original source name', after_src == before_src, (before_src[:3], after_src[:3]))
    if not LIVE:
        pg.evaluate("AutoprodLicence.activate('TEST-PRO-YEARLY')"); pg.wait_for_timeout(300); pg.evaluate('renderApp()')
        pg.fill('#proTemplateName', 'Awkward test'); pg.click('#proSaveTemplateBtn')
        up(pg, S + 'sample-recipes.pdf'); wait_status(pg, 'Sorted')
        pg.evaluate("loadTemplate(encodeURIComponent('Awkward test'))"); pg.wait_for_timeout(300)
        check('template: loading a saved template keeps original source names', pg.evaluate('ROWS.map(r=>r.srcSheet)') == before_src)
        up(pg, S + 'sample-recipes.pdf'); wait_status(pg, 'Sorted')
        pg.evaluate("document.getElementById('statusMsg').textContent=''"); pg.set_input_files('#appendInput', S + 'fixes-and-titles.pdf'); wait_status(pg, 'Appended|Could not', 30000)
        srcs = pg.evaluate('[...new Set(ROWS.map(r=>r.srcSheet))]')
        check('append PDF: old and new rows keep their own recipe titles', srcs == ['Roast Pumpkin Salad', 'Braised Lamb Shoulder', 'Braised Lamb Shoulder: 18 Hour [GF]', "Market Fish's Salsa Verde"], srcs)
        pg.wait_for_timeout(1000); pg.reload(); pg.wait_for_load_state('networkidle'); pg.wait_for_timeout(600)
        check('append + reload: source names and PDF totals survive', pg.evaluate('[...new Set(ROWS.map(r=>r.srcSheet))]') == srcs and pg.evaluate('ROWS.every(r=>r.pdfFixedTotal!==undefined)'))
        pg.evaluate("AutoprodLicence.deactivate()"); pg.wait_for_timeout(300)
    # Source tab column screenshot: render the section export like a sheet
    rows, _ = rows_of(os.path.join(OUT, 'awkward-by-section.xlsx'))
    tbl = ''.join(f'<tr><td>{t}</td><td>{i}</td><td><b>{sv.replace(" ", "&middot;")}</b></td></tr>' for t, i, sv in rows)
    sp = ctx.new_page(); sp.set_viewport_size({'width': 900, 'height': 700})
    sp.set_content(f'<style>body{{font:13px sans-serif;padding:16px}}td,th{{border:1px solid #bbb;padding:4px 8px}}table{{border-collapse:collapse}}</style><h3>awkward-by-section.xlsx, opened with openpyxl (spaces shown as &middot;)</h3><table><tr><th>Tab (sheet)</th><th>Item</th><th>Source tab column</th></tr>{tbl}</table>')
    sp.screenshot(path=f'{SHOTS}/06-export-source-tab-column.png', full_page=True); sp.close()
    # free multi-PDF (limits stay off)
    up(pg, [S + 'sample-recipes.pdf', '/home/box/prototypes/monetise/tests/samples/sample-prep-B.pdf']); st5 = wait_status(pg, 'Sorted|Could not', 30000)
    check('free plan: several PDFs at once still allowed (free limits off)', len(set(pg.evaluate('ROWS.map(r=>r.dish)'))) == 3, st5[:120])
    check('free plan: no "Made with Prepline" footer on free prints', 'Made with Prepline' not in capture_print(pg, "printAllSections()"))
    # ---------- G6. v3.2.2: pax changes rescale everything, live ----------
    def vrow(task, view):
        pg.evaluate(f"setView('{view}')"); pg.wait_for_timeout(120)
        return ' | '.join(pg.locator(f'#resultsArea .section-table-wrap.no-print tr:has(td:has-text("{task}"))').all_inner_texts())
    def pax_box(loc, dish):
        enc = pg.evaluate("d => encName(d)", dish)
        return pg.locator(f'input[data-pax-loc="{loc}"][data-pax-dish="{enc}"]')
    def type_pax(loc, dish, text, view='dish'):
        pg.evaluate(f"setView('{view}')"); pg.wait_for_timeout(120)
        box = pax_box(loc, dish); box.click(); pg.keyboard.press('Control+A'); pg.keyboard.press('Backspace')
        pg.keyboard.type(text, delay=30); pg.wait_for_timeout(150)   # no blur, no Enter: must already be live
        return pg.evaluate("[document.activeElement.dataset.paxLoc, decodeURIComponent(document.activeElement.dataset.paxDish||''), document.activeElement.value]")
    def blur(): pg.evaluate("document.activeElement && document.activeElement.blur()"); pg.wait_for_timeout(150)
    def xlsx_totals(path):
        f = openpyxl.load_workbook(path); v = openpyxl.load_workbook(path, data_only=True); out = {}
        for ws in f.worksheets:
            wv = v[ws.title]; pax = None; hdr = False
            for (rf, rv) in zip(ws.iter_rows(), wv.iter_rows()):
                a = rf[0].value
                if a == 'PAX': pax = rf[1].value; hdr = False; continue
                if a == 'Item': hdr = True; continue
                if a is None: continue
                if hdr and any(c.value is not None for c in rf[1:5]): out[(ws.title, a, pax)] = (rf[4].value, rv[4].value, rf[2].value)
        return out
    L, BF = 'Lamb Rump Sat', 'Beef Fillet Sat'
    def txt(h): return re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', re.sub(r'<style.*?</style>', '', h, flags=re.S)))
    errs.clear(); bad_urls = []; pg.on('response', lambda r: r.status >= 400 and bad_urls.append(f'{r.status} {r.url[:120]}'))
    up(pg, S + 'pax-scaling.xlsx'); wait_status(pg, 'Sorted')
    g0 = vrow('Garlic confit', 'dish')
    foc = type_pax('dish', L, '15')
    g1, r1 = vrow('Garlic confit', 'dish'), vrow('Lamb rump trim', 'dish')
    check('pax live (By dish): typing 15 rescales while typing, no blur needed: 90 g x 10 = 900 g → 90 g x 15 = 1.35 kg; 0.2 kg x 15 = 3 kg', '900 g' in g0 and '1.35 kg' in g1 and '3 kg' in r1, (g0, g1, r1))
    check('pax live: cursor stays in the pax box with what was typed', foc == ['dish', L, '15'], foc)
    s_on, s_r = vrow('Onion dice', 'section'), vrow('Lamb rump trim', 'section')
    check('pax live (By team section): 3 kg row + combined Onion dice 0.09x15 + 0.05x20 = 2.35 kg', '3 kg' in s_r and '2.35 kg' in s_on and '×2' in s_on, (s_r, s_on))
    d_g, d_on = vrow('Garlic confit', 'priority'), vrow('Onion dice', 'priority')
    check('pax live (By prep day): 1.35 kg row + combined Onion dice 2.35 kg', '1.35 kg' in d_g and '2.35 kg' in d_on, (d_g, d_on))
    foc = type_pax('panel', BF, '30', view='section')
    s_on, s_b = vrow('Onion dice', 'section'), vrow('Beef fillet portion', 'section')
    check('pax panel box is live too (Beef 30): beef 5.4 kg, combined Onion dice 1.35 + 1.5 = 2.85 kg', '5.4 kg' in s_b and '2.85 kg' in s_on and foc == ['panel', BF, '30'], (s_b, s_on, foc))
    pg.evaluate("setView('dish')"); pg.wait_for_timeout(150)
    check('pax boxes stay in step: the By dish box for Beef now shows 30', pax_box('dish', BF).input_value() == '30', pax_box('dish', BF).input_value())
    type_pax('dish', L, '12.5'); w, rj = vrow('Lamb rump trim', 'dish'), vrow('Red wine jus', 'dish')
    check('decimal pax 12.5: 0.2 kg → 2.5 kg; 0.05 L → 625 ml (L rolls back to ml)', '2.5 kg' in w and '625 ml' in rj, (w, rj))
    type_pax('dish', L, '15'); blur()
    # prints
    hd = txt(capture_print(pg, f"printDish(encodeURIComponent({json.dumps(L)}))"))
    hs, hy = txt(capture_print(pg, "printAllSections()")), txt(capture_print(pg, "printAllDays()"))
    pg.evaluate("setView('section')"); hsec = txt(capture_print(pg, "printSection('Vegetables')"))
    pg.evaluate("setView('priority')"); hday = txt(capture_print(pg, "printPriority([...document.querySelectorAll('[id^=print-priority-]')].map(e=>e.id.replace('print-priority-','')).find(k=>k.endsWith('Wednesday')))"))
    check('print dish sheet: PAX 15 and 1.35 kg / 3 kg', 'PAX / YIELD (COVERS): 15 ' in hd and '1.35 kg' in hd and '3 kg' in hd, len(hd))
    check('print section + day handouts: combined Onion dice 2.85 kg', '2.85 kg' in hsec and '2.85 kg' in hday, (len(hsec), len(hday)))
    check('print all sections / all days (one page complete): 1.35 kg, 2.85 kg, 5.4 kg', all(x in hs and x in hy for x in ['1.35 kg', '2.85 kg', '5.4 kg']), (len(hs), len(hy)))
    # exports
    want = {('Lamb rump trim', 15): 3, ('Garlic confit', 15): 1350, ('Onion dice', 15): 1.35, ('Red wine jus', 15): 0.75, ('Beef fillet portion', 30): 5.4, ('Onion dice', 30): 1.5}
    for view, fn in [('dish', 'pax-by-dish.xlsx'), ('section', 'pax-by-section.xlsx'), ('priority', 'pax-by-day.xlsx')]:
        path, _, _ = export_view(pg, view, fn); t = xlsx_totals(path)
        got = {(it, px): val for (tab, it, px), (f, val, u) in t.items()}
        forms = [f for (f, val, u) in t.values()]
        bad = {k: (got.get(k), v) for k, v in want.items() if got.get(k) is None or abs(got[k] - v) > 1e-9}
        check(f'export {view} (openpyxl): PAX cells 15/30 and scaled Total needed values', not bad and all(isinstance(f, str) and re.fullmatch(r'=B(\d+)\*D\1', f) for f in forms), bad or forms[:3])
    # reload
    pg.wait_for_timeout(1000); pg.reload(); pg.wait_for_load_state('networkidle'); pg.wait_for_timeout(600)
    px = pg.evaluate('paxByDish'); g2 = vrow('Garlic confit', 'dish'); o2 = vrow('Onion dice', 'priority')
    check('reload: edited pax (15, 30) and the scaled totals survive', px == {L: 15, BF: 30} and '1.35 kg' in g2 and '2.85 kg' in o2, (px, g2, o2))
    # edge cases
    type_pax('dish', L, ''); e0 = vrow('Garlic confit', 'dish')
    check('blank pax = not set: spreadsheet total says "set pax"', 'set pax' in e0 and L not in pg.evaluate('paxByDish'), e0)
    type_pax('dish', L, '0'); e1 = vrow('Garlic confit', 'dish')
    check('pax 0: totals are 0 (not "set pax", no NaN)', '0 g' in e1 and 'NaN' not in e1, e1)
    pg.evaluate("setView('dish')"); bx = pax_box('dish', L); bx.click(); pg.keyboard.press('Control+A'); pg.keyboard.type('abc', delay=30); pg.wait_for_timeout(150)
    bad_cls = pax_box('dish', L).get_attribute('class') or ''
    check('junk pax "abc" refused: box marked red, last good pax (0) kept', 'pax-invalid' in bad_cls and pg.evaluate(f'paxByDish[{json.dumps(L)}]') == 0, (bad_cls, pg.evaluate('paxByDish')))
    blur(); check('junk pax: leaving the box explains it', 'must be a number' in pg.inner_text('#statusMsg'), pg.inner_text('#statusMsg'))
    type_pax('dish', L, '1,000,000'); e2 = vrow('Garlic confit', 'dish')
    check('very large pax 1,000,000: 90 g x 1,000,000 = 90000 kg, no "e+" or NaN', '90000 kg' in e2 and 'e+' not in e2 and 'NaN' not in e2, e2)
    type_pax('dish', L, '2000000')
    check('pax above 1,000,000 refused: box marked red, 2,000,000 never stored', pg.evaluate(f'paxByDish[{json.dumps(L)}]') <= 1000000 and 'pax-invalid' in (pax_box('dish', L).get_attribute('class') or ''), pg.evaluate('paxByDish'))
    type_pax('dish', L, '15'); blur()
    # PDF dishes: scale from the front-page yield
    up(pg, S + 'fixes-and-titles.pdf'); wait_status(pg, 'Sorted|Could not', 30000)
    MF, BL = "Market Fish's Salsa Verde", 'Braised Lamb Shoulder: 18 Hour [GF]'
    base = pg.evaluate("Object.fromEntries(ROWS.map(r=>[r.dish, r.pdfBasePax]))")
    b0 = vrow('Butternut pumpkin', 'dish')
    check('PDF: base pax = front-page Yield (60 and 100) and totals start as printed (900 g)', base == {BL: 100, MF: 60} and '900 g' in b0, (base, b0))
    type_pax('dish', MF, '90')
    pb, pc, pz, pl = vrow('Butternut pumpkin', 'dish'), vrow('Chicken thigh', 'dish'), vrow('Lemon zest', 'dish'), vrow('Lamb shoulder', 'dish')
    check('PDF pax 60 → 90 (x1.5): 900 g → 1.35 kg, 2.5 kg → 3.75 kg, 200 g → 300 g; other dish unchanged (12 kg)', '1.35 kg' in pb and '3.75 kg' in pc and '300 g' in pz and '12 kg' in pl, (pb, pc, pz, pl))
    check('PDF: section and day views show the scaled 1.35 kg', '1.35 kg' in vrow('Butternut pumpkin', 'section') and '1.35 kg' in vrow('Butternut pumpkin', 'priority'))
    check('PDF: print all days has 1.35 kg', '1.35 kg' in txt(capture_print(pg, 'printAllDays()')))
    path, _, _ = export_view(pg, 'dish', 'pax-pdf-by-dish.xlsx'); t = xlsx_totals(path)
    pv = {it: (f, val) for (tab, it, px), (f, val, u) in t.items() if px == 90}
    check('PDF export (openpyxl): PAX 90, Total needed = per-pax x PAX in g (1350 g, 3750 g)', abs(pv.get('Butternut pumpkin', (0, 0))[1] - 1350) < 1e-6 and abs(pv.get('Chicken thigh', (0, 0))[1] - 3750) < 1e-6 and str(pv.get('Butternut pumpkin', ('',))[0]).startswith('=B'), pv)
    pg.wait_for_timeout(1000); pg.reload(); pg.wait_for_load_state('networkidle'); pg.wait_for_timeout(600)
    rb = vrow('Butternut pumpkin', 'dish'); rtxt = pg.inner_text('#resultsArea')
    check('PDF reload: pax 90 and 1.35 kg survive; still no "set pax" on PDF rows (v3.2 fix kept)', pg.evaluate(f'paxByDish[{json.dumps(MF)}]') == 90 and '1.35 kg' in rb and 'set pax' not in rtxt, rb)
    type_pax('dish', MF, ''); check('PDF blank pax: back to the printed 900 g', '900 g' in vrow('Butternut pumpkin', 'dish'))
    # a list saved by v3.2.1 or earlier (no base yield stored) still scales after reload
    pg.evaluate(f"""(()=>{{const k='preplistOnetimeSession_v1'; const s=JSON.parse(localStorage.getItem(k)); s.ROWS.forEach(r=>delete r.pdfBasePax); s.paxByDish[{json.dumps(MF)}]=60; localStorage.setItem(k, JSON.stringify(s));}})()""")
    pg.reload(); pg.wait_for_load_state('networkidle'); pg.wait_for_timeout(600)
    mg = pg.evaluate("Object.fromEntries(ROWS.map(r=>[r.dish, r.pdfBasePax]))"); type_pax('dish', MF, '90')
    check('older save without a stored yield: base restored from its pax (60), then 90 pax → 1.35 kg', mg.get(MF) == 60 and '1.35 kg' in vrow('Butternut pumpkin', 'dish'), mg)
    up(pg, S + 'no-yield.pdf'); wait_status(pg, 'Sorted|Could not', 30000)
    ny = 'Herb Focaccia'; type_pax('dish', ny, '50'); nyr = vrow('Bread flour', 'dish')
    check('PDF with no Yield line: totals stay as printed (3 kg) and the pax panel says why', '3 kg' in nyr and 'no readable Yield' in pg.inner_text('#paxList'), (nyr, pg.inner_text('#paxList')))
    # append + templates (Pro)
    if not LIVE:
        pg.evaluate("AutoprodLicence.activate('TEST-PRO-YEARLY')"); pg.wait_for_timeout(300); pg.evaluate('renderApp()')
        up(pg, S + 'sample-recipes.pdf'); wait_status(pg, 'Sorted')
        pg.evaluate("document.getElementById('statusMsg').textContent=''"); pg.set_input_files('#appendInput', S + 'fixes-and-titles.pdf'); wait_status(pg, 'Appended|Could not', 30000)
        type_pax('dish', MF, '90'); type_pax('dish', 'Roast Pumpkin Salad', '60')
        ab = vrow('Butternut pumpkin', 'dish')
        check('append PDF: appended dish scales (1.35 kg) and the first PDF scales on its own yield (120 → 60: 4 kg → 2 kg)', '1.35 kg' in ab and '2 kg' in ab and pg.evaluate("ROWS.filter(r=>r.dish==='Braised Lamb Shoulder')[0].pdfBasePax") == 120, ab)
        blur(); pg.wait_for_timeout(1000); pg.reload(); pg.wait_for_load_state('networkidle'); pg.wait_for_timeout(600)
        check('append + reload: both pax edits and scaled totals survive', '1.35 kg' in vrow('Butternut pumpkin', 'dish') and '2 kg' in vrow('Butternut pumpkin', 'dish'))
        pg.fill('#proTemplateName', 'Pax test'); pg.click('#proSaveTemplateBtn')
        up(pg, S + 'pax-scaling.xlsx'); wait_status(pg, 'Sorted')
        pg.evaluate("loadTemplate(encodeURIComponent('Pax test'))"); pg.wait_for_timeout(300)
        tb = vrow('Butternut pumpkin', 'dish'); type_pax('dish', MF, '120'); tb2 = vrow('Butternut pumpkin', 'dish')
        check('template: reloads pax 90/60 with scaled totals, and still rescales live (120 → 1.8 kg)', '1.35 kg' in tb and '2 kg' in tb and '1.8 kg' in tb2, (tb, tb2))
        pg.evaluate("AutoprodLicence.deactivate()"); pg.wait_for_timeout(300)
    # recipe database (weekly) dishes
    pg.evaluate("setToolMode('database')"); up(pg, S + 'pax-scaling.xlsx'); wait_status(pg, 'Added|Could not')
    up(pg, S + 'fixes-and-titles.pdf'); wait_status(pg, 'Added|Could not', 30000)
    pg.evaluate("appMode='weekly'; renderApp()")
    pg.fill(f'#weeklyDishList tr:has-text("{L}") input[placeholder="pax"]', '15'); pg.evaluate('generateWeekly()'); pg.wait_for_timeout(300)
    dg, dbf = vrow('Garlic confit', 'dish'), vrow('Butternut pumpkin', 'dish')
    check('recipe database: weekly pax 15 → 1.35 kg; a PDF dish with no pax picked starts at its yield (900 g)', '1.35 kg' in dg and '900 g' in dbf, (dg, dbf))
    type_pax('dish', L, '20'); dg2 = vrow('Garlic confit', 'dish')
    check('recipe database: editing pax in the list rescales live (1.8 kg) and updates this week\'s pick', '1.8 kg' in dg2 and pg.evaluate(f'weeklySelection[{json.dumps(L)}].pax') == '20', dg2)
    wbox = pg.locator(f'#weeklyDishList tr:has-text("{L}") input[placeholder="pax"]'); wbox.click(); pg.keyboard.press('Control+A'); pg.keyboard.type('25', delay=30); pg.wait_for_timeout(150)
    type_pax('dish', MF, '90'); dg3, dbf3 = vrow('Garlic confit', 'dish'), vrow('Butternut pumpkin', 'dish')
    check('recipe database: weekly pax box is live after Generate (25 → 2.25 kg); database PDF dish scales (1.35 kg)', '2.25 kg' in dg3 and '1.35 kg' in dbf3, (dg3, dbf3))
    pg.evaluate("dishDatabase={}; weeklySelection={}; saveDatabase(); setToolMode('onetime')")
    check('v3.2.2 pax checks: no JS errors', not [e for e in errs if 'Failed to load resource' not in e], (errs, bad_urls))
    results.append({'test': 'v3.2.2 failed loads (info)', 'ok': True, 'info': str(bad_urls)})
    # ---------- G7. v3.2.3: Quantity per pax column + live Excel formulas ----------
    import shutil, subprocess
    HDR = ['Item', 'Quantity per pax', 'Unit', 'PAX (this row)', 'Total needed', 'Section', 'Prep day', 'Source tab']
    REF = re.compile(r'\$?[A-Z]{1,3}\$?\d+')
    HAVE_LO = bool(shutil.which('soffice'))
    def sheet_rows(path):
        wb = openpyxl.load_workbook(path); out = []
        for ws in wb.worksheets:
            paxrow, hdr = None, None
            for row in ws.iter_rows(max_col=8):
                a = row[0].value
                if a == 'PAX': paxrow, hdr = row[0].row, None; continue
                if a == 'Item': hdr = [c.value for c in row]; continue
                if a is None or not hdr or all(c.value is None for c in row[1:]): continue
                out.append({'tab': ws.title, 'r': row[0].row, 'paxrow': paxrow, 'hdr': hdr, 'item': a, 'B': row[1].value, 'C': row[2].value,
                            'D': row[3].value, 'E': row[4].value, 'src': row[7].value})
        return out
    def py_eval(path):   # fallback evaluator for the export's own formula shapes, if LibreOffice is missing
        wb = openpyxl.load_workbook(path)
        def val(ws, ref, depth=0):
            v = ws[ref.replace('$', '')].value
            if isinstance(v, str) and v.startswith('='): return ev(ws, v[1:], depth + 1)
            return '' if v is None else v
        def ev(ws, f, depth):
            m = re.fullmatch(r'IF\((.+?)="","",(.+)\)', f)
            if m: a = val(ws, m.group(1)); return '' if a == '' else ev(ws, m.group(2), depth)
            m = re.fullmatch(r'IF\(ISNUMBER\((\w+)\),(.+),""\)', f)
            if m: return ev(ws, m.group(2), depth) if isinstance(val(ws, m.group(1)), (int, float)) else ''
            m = re.fullmatch(r'(\$?\w+\$?\d+)\*(\$?\w+\$?\d+)', f)
            if m:
                a, c = val(ws, m.group(1)), val(ws, m.group(2))
                return a * c if isinstance(a, (int, float)) and isinstance(c, (int, float)) else ''
            return val(ws, f, depth)
        res = {}
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_col=5):
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith('='): res[(ws.title, c.coordinate)] = ev(ws, c.value[1:], 0)
        return res
    def recalc(paths):
        """Recalculated values {path: {(tab, cell): value}} — LibreOffice headless if present."""
        out = {}
        if HAVE_LO:
            d = os.path.join(OUT, 'lo-recalc'); os.makedirs(d, exist_ok=True)
            subprocess.run(['soffice', '--headless', '--norestore', '--convert-to', 'xlsx', '--outdir', d] + paths, capture_output=True, timeout=300)
            for pth in paths:
                wb = openpyxl.load_workbook(os.path.join(d, os.path.basename(pth)), data_only=True)
                out[pth] = {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws.iter_rows(max_col=5) for c in row}
        else:
            for pth in paths: out[pth] = py_eval(pth)
        return out
    def app_vals():
        return {(x[0], x[1]): x[2:] for x in pg.evaluate("""activeRows().map(r=>{const t=computeTotal(r); const b=toBaseQty(t, r.unit);
            return [srcName(r), r.task, t, r.unit||'', b.n, b.u, r.pdfFixedTotal!=null, Number.isFinite(r.pdfBasePax)&&r.pdfBasePax>0, r.dish]})""")}
    def app_in(unit, v):
        t, u, bn, bu = v[0], v[1], v[2], v[3]
        if t is None: return None
        return t if (unit or '') == u else (bn if unit == bu else ('unit mismatch', unit, u))
    def close(a, b2): return isinstance(a, (int, float)) and isinstance(b2, (int, float)) and abs(a - b2) <= 1e-6 * max(1, abs(b2))
    def g7(label):
        vals0 = app_vals(); files = []
        for view in ['dish', 'section', 'priority']:
            path, _, _ = export_view(pg, view, f'g7-{label}-{view}.xlsx'); files.append(path)
        rows_all = [sheet_rows(f) for f in files]; bad = []
        for rows in rows_all:
            for x in rows:
                v = vals0.get((x['src'], x['item']))
                if x['hdr'] != HDR: bad.append(('header', x['hdr']))
                if v is None: bad.append(('no app row', x['src'], x['item'])); continue
                if v[4] and not v[5]:
                    if not (x['B'] is None and x['D'] is None and isinstance(x['E'], (int, float))): bad.append(('no-yield row', x))
                    continue
                if not (isinstance(x['D'], str) and x['D'].startswith('=') and f"$B${x['paxrow']}" in x['D']): bad.append(('PAX cell', x))
                if not (isinstance(x['E'], str) and x['E'].startswith('=') and f"B{x['r']}" in x['E'] and f"D{x['r']}" in x['E']): bad.append(('Total cell', x))
                for f in (x['D'], x['E']):
                    if isinstance(f, str) and re.search(r'\d', REF.sub('', f)): bad.append(('number inside formula', f))
                if v[4] and v[1] in ('kg', 'g', 'L', 'ml') and x['C'] not in ('g', 'ml'): bad.append(('PDF unit not g/ml', x['item'], x['C']))
        check(f'{label}: all 3 exports — every row has Quantity per pax + PAX (this row); Total needed = per-pax cell x PAX cell (cell refs only, no numbers in formulas)',
              not bad and all(rows_all), bad[:3])
        # edit in "Excel": every PAX +7; in the section file also double one row's per-pax and override one row's PAX
        edited, special = [], {}
        for f, rows in zip(files, rows_all):
            wb = openpyxl.load_workbook(f)
            for ws in wb.worksheets:
                for row in ws.iter_rows(max_col=2):
                    if row[0].value == 'PAX' and isinstance(row[1].value, (int, float)): row[1].value = row[1].value + 7
            if f.endswith('-section.xlsx'):
                fr = [x for x in rows if isinstance(x['E'], str)]
                if len(fr) >= 2:
                    a, c = fr[0], fr[1]
                    wb[a['tab']][f"B{a['r']}"] = a['B'] * 2; wb[c['tab']][f"D{c['r']}"] = 3
                    special = {(a['tab'], f"E{a['r']}"): ('x2', a), (c['tab'], f"E{c['r']}"): ('pax3', c)}
            e = f.replace('.xlsx', '-edited.xlsx'); wb.save(e); edited.append(e)
        rc = recalc(files + edited)
        for d0, px in pg.evaluate("Object.entries(paxByDish)"): pg.evaluate(f"setPax(encName({json.dumps(d0)}), '{px + 7}')")
        vals1 = app_vals(); pg.wait_for_timeout(100)
        mism0, mism1 = [], []
        for f, e, rows in zip(files, edited, rows_all):
            for x in rows:
                cell = (x['tab'], f"E{x['r']}")
                v0, v1 = vals0[(x['src'], x['item'])], vals1[(x['src'], x['item'])]
                exp0, got0 = app_in(x['C'], v0), rc[f].get(cell)
                if not ((exp0 is None and got0 in (None, '')) or close(got0, exp0)): mism0.append((x['tab'], x['item'], got0, exp0))
                if cell in special and e.endswith('-section-edited.xlsx'):
                    kind, a = special[cell]; pax1 = rc[e].get((x['tab'], f"D{x['r']}"))
                    exp1 = a['B'] * 2 * pax1 if kind == 'x2' else a['B'] * 3
                else: exp1 = app_in(x['C'], v1)
                got1 = rc[e].get(cell)
                if not ((exp1 is None and got1 in (None, '')) or close(got1, exp1)): mism1.append((x['tab'], x['item'], got1, exp1))
        how = 'LibreOffice headless' if HAVE_LO else 'Python formula evaluator'
        check(f'{label}: recalculated ({how}) totals match the app exactly as exported', not mism0, mism0[:3])
        check(f'{label}: after editing PAX (+7) in Excel, recalculated totals match the app with the same pax; edited per-pax (x2) and one row\'s own PAX (3) recalc just that row',
              not mism1 and (len(special) == 2 or label.startswith('PDF no Yield')), (mism1[:3], len(special)))
        return files
    # spreadsheet
    up(pg, S + 'pax-scaling.xlsx'); wait_status(pg, 'Sorted'); g7('spreadsheet')
    # PDF with Yield: per-pax in the app too
    up(pg, S + 'fixes-and-titles.pdf'); wait_status(pg, 'Sorted|Could not', 30000)
    pp = {t: vrow(t, 'dish') for t in ['Chicken thigh', 'Butternut pumpkin', 'Lemon zest', 'Lamb shoulder', 'Red wine']}
    check('app Per pax (PDF, printed total ÷ Yield): 2.5 kg/60 = 41.7 g, 900 g/60 = 15 g, 200 g/60 = 3.33 g, 12 kg/100 = 120 g, 1.5 L/100 = 15 ml',
          '41.7 g' in pp['Chicken thigh'] and '15 g' in pp['Butternut pumpkin'] and '3.33 g' in pp['Lemon zest'] and '120 g' in pp['Lamb shoulder'] and '15 ml' in pp['Red wine'], pp)
    pg.evaluate('setConsolidate(false)'); sv = vrow('Chicken thigh', 'section'); pg.evaluate("setConsolidate(true); setView('dish')"); hp = txt(capture_print(pg, f"printDish(encodeURIComponent({json.dumps(MF)}))"))
    check('app Per pax shows in By team section and on the printed dish sheet too', '41.7 g' in sv and '41.7 g' in hp, sv)
    pfiles = g7('PDF')
    up(pg, pfiles[0]); wait_status(pg, 'Sorted|Could not')
    rt = vrow('Chicken thigh', 'dish')
    check('round trip: re-uploading the PDF By dish export keeps the total (41.67 g x 60 = 2.5 kg)', '2.5 kg' in rt, rt)
    # PDF with no Yield
    up(pg, S + 'no-yield.pdf'); wait_status(pg, 'Sorted|Could not', 30000)
    nyp = vrow('Bread flour', 'dish')
    nf = g7('PDF no Yield'); nr = {x['item']: x for x in sheet_rows(nf[0])}
    check('PDF no Yield: app Per pax is "—"; export leaves Quantity per pax and PAX blank and writes the printed total (3 kg, 800 g) as a plain value',
          '—' in nyp and nr['Bread flour']['B'] is None and nr['Bread flour']['E'] == 3 and nr['Bread flour']['C'] == 'kg' and nr['Rosemary oil']['E'] == 800, (nyp, nr))
    # append + template (Pro)
    if not LIVE:
        pg.evaluate("AutoprodLicence.activate('TEST-PRO-YEARLY')"); pg.wait_for_timeout(300); pg.evaluate('renderApp()')
        up(pg, S + 'sample-recipes.pdf'); wait_status(pg, 'Sorted')
        pg.evaluate("document.getElementById('statusMsg').textContent=''"); pg.set_input_files('#appendInput', [S + 'fixes-and-titles.pdf', S + 'no-yield.pdf']); wait_status(pg, 'Appended|Could not', 30000)
        pg.evaluate(f"setPax(encName({json.dumps(MF)}), '90')"); g7('append')
        pg.fill('#proTemplateName', 'Per pax test'); pg.click('#proSaveTemplateBtn')
        up(pg, S + 'pax-scaling.xlsx'); wait_status(pg, 'Sorted')
        pg.evaluate("loadTemplate(encodeURIComponent('Per pax test'))"); pg.wait_for_timeout(300); g7('template')
        pg.evaluate("AutoprodLicence.deactivate()"); pg.wait_for_timeout(300)
    check('v3.2.3 per-pax checks: no JS errors', not [e for e in errs if 'Failed to load resource' not in e], errs)
    # ---------- G8. v3.2.4: "Include method in printout" Yes/No ----------
    ctx8, _ = new_ctx(b); pg, errs = page_with_errs(ctx8)
    pg.goto(URL); pg.wait_for_load_state('networkidle')
    up(pg, S + 'sample-recipes.pdf'); wait_status(pg, 'Sorted')
    def pmt_state(): return pg.evaluate("[includeMethodInPrint, localStorage.getItem('preplinePref_printIncludeMethod_v1'), document.querySelector('#printMethodToggle [data-pmt=yes]').getAttribute('aria-pressed')]")
    beside = []
    for v in ['dish', 'section', 'priority']:
        pg.evaluate(f"setView('{v}')"); pg.wait_for_timeout(150)
        beside.append(pg.evaluate("(()=>{const t=document.getElementById('printMethodToggle'); const row=t&&t.closest('.print-row'); return !!(row && row.querySelector('button[onclick^=printAllSections]') && t.offsetParent && t.innerText.includes('Include method in printout'));})()"))
    check('print toggle: "Include method in printout" Yes/No sits in the print-button row in all 3 views', beside == [True, True, True], beside)
    check('print toggle: default is No (nothing stored yet), Yes not pressed', pmt_state() == [False, None, 'false'] and pg.evaluate("document.querySelector('#printMethodToggle [data-pmt=no]').getAttribute('aria-pressed')") == 'true', pmt_state())
    def all_prints():
        out = {}
        pg.evaluate("setView('dish')"); pg.wait_for_timeout(150)
        out['dish sheet'] = capture_print(pg, "printDish(encodeURIComponent(Object.keys(methodCrops)[0]))")
        pg.evaluate("setView('section')"); pg.wait_for_timeout(150)
        out['section handout'] = capture_print(pg, "printSection(ROWS[0].section)")
        pg.evaluate("setView('priority')"); pg.wait_for_timeout(150)
        out['prep-day handout'] = capture_print(pg, "printPriority([...document.querySelectorAll('[id^=print-priority-]')].map(e=>e.id.replace('print-priority-',''))[0])")
        out['all sections (one page complete)'] = capture_print(pg, 'printAllSections()')
        out['all prep days (one page complete)'] = capture_print(pg, 'printAllDays()')
        pg.evaluate("setView('dish')"); pg.wait_for_timeout(150)
        return out
    def nostyle(h): return re.sub(r'<style.*?</style>', '', h, flags=re.S)
    def has_method(h): h = nostyle(h); return any(x in h for x in ['Method reference', 'orig-pages-grid"', 'class="orig-page-item', '<img', 'data:image'])
    dflt = all_prints()
    check('v3.2.5 default No: every print type (dish, section, prep day, both one-page-complete) has no method captures', not any(has_method(h) for h in dflt.values()) and all('<table' in h for h in dflt.values()), {k: has_method(v) for k, v in dflt.items()})
    pg.click('#printMethodToggle [data-pmt=yes]'); pg.wait_for_timeout(200)
    check('Yes: toggle shows Yes pressed and stores yes', pmt_state() == [True, 'yes', 'true'], pmt_state())
    yes = all_prints()
    check('Yes: every print type (dish, section, prep day, both one-page-complete) includes the Method captures', all(has_method(h) and 'Method reference' in h for h in yes.values()), {k: has_method(v) for k, v in yes.items()})
    pg.click('#printMethodToggle [data-pmt=no]'); pg.wait_for_timeout(200)
    check('No: toggle shows No pressed and is stored under its own key', pmt_state() == [False, 'no', 'false'] and 'preplinePref' not in (pg.evaluate("localStorage.getItem('preplistOnetimeSession_v1')") or ''), pmt_state())
    no = all_prints()
    for k, h in no.items():
        check(f'No: {k} print has no method images, no "Method reference" heading, no empty method block', not has_method(h) and 'method-crops-title' not in nostyle(h) and 'orig-pages-grid' not in nostyle(h) and '<table' in h, len(h))
    dom = pg.evaluate("[document.querySelectorAll('#resultsArea .print-only .orig-pages-grid, #resultsArea .print-only img').length, document.querySelectorAll('button').length && [...document.querySelectorAll('button')].filter(b=>b.innerText.includes('View original')).length]")
    check('No: print DOM has no method blocks; on-screen 📖 View original buttons stay', dom[0] == 0 and dom[1] == 2, dom)
    pp = ctx8.new_page(); pp.set_viewport_size({'width': 900, 'height': 1200}); pp.set_content(no['dish sheet']); pp.emulate_media(media='print'); pp.wait_for_timeout(500)
    gap = pp.evaluate("(()=>{const t=[...document.querySelectorAll('.print-only table')].pop(); const r=t.getBoundingClientRect(); return [Math.round(document.body.scrollHeight - (r.bottom + window.scrollY)), document.querySelectorAll('img').length];})()")
    pp.screenshot(path=f'{SHOTS}/09-print-dish-method-off.png', full_page=True); pp.close()
    check('No: printed dish sheet closes up after the table (no blank method page)', gap[1] == 0 and gap[0] < 150, gap)
    pg.wait_for_timeout(600); pg.reload(); pg.wait_for_load_state('networkidle'); pg.wait_for_timeout(600)
    check('No survives a reload (toggle still No, prints still without method)', pmt_state() == [False, 'no', 'false'] and not has_method(capture_print(pg, "printAllDays()")), pmt_state())
    pg.click('#printMethodToggle [data-pmt=yes]'); pg.wait_for_timeout(200); pg.reload(); pg.wait_for_load_state('networkidle'); pg.wait_for_timeout(600)
    check('Yes survives a reload too, and method captures come back', pmt_state() == [True, 'yes', 'true'] and 'Method reference' in capture_print(pg, "printAllSections()"), pmt_state())
    pg.evaluate("document.getElementById('printMethodToggle').scrollIntoView({block:'center'})"); pg.wait_for_timeout(200)
    pg.screenshot(path=f'{SHOTS}/08-print-method-toggle-desktop.png')
    check('print toggle: no JS errors', not errs, errs)
    ctx8.close()
    # v3.2.5: a 'yes' remembered from before (v3.2.4 default) still prints the method; 'no' stays No
    for stored, want in [('yes', True), ('no', False)]:
        ctx8, _ = new_ctx(b); pg, errs = page_with_errs(ctx8)
        pg.goto(URL); pg.wait_for_load_state('networkidle')
        pg.evaluate(f"localStorage.setItem('preplinePref_printIncludeMethod_v1', '{stored}')"); pg.reload(); pg.wait_for_load_state('networkidle')
        up(pg, S + 'sample-recipes.pdf'); wait_status(pg, 'Sorted')
        hs = capture_print(pg, "printDish(encodeURIComponent(Object.keys(methodCrops)[0]))")
        check(f"v3.2.5: a stored '{stored}' stays {'Yes' if want else 'No'} (toggle and dish print {'with' if want else 'without'} method)", pmt_state() == [want, stored, 'true' if want else 'false'] and has_method(hs) == want and ('Method reference' in hs) == want and not errs, (pmt_state(), errs))
        pg.wait_for_timeout(1000); pg.reload(); pg.wait_for_load_state('networkidle'); pg.wait_for_function("document.getElementById('printMethodToggle')", timeout=10000)
        check(f"v3.2.5: stored '{stored}' survives a reload", pmt_state() == [want, stored, 'true' if want else 'false'], pmt_state())
        ctx8.close()
    ctx8, _ = new_ctx(b, 390, 844, is_mobile=True, has_touch=True, device_scale_factor=2); pg, errs = page_with_errs(ctx8)
    pg.goto(URL); pg.wait_for_load_state('networkidle'); up(pg, S + 'sample-recipes.pdf'); wait_status(pg, 'Sorted')
    box = pg.evaluate("(()=>{const r=document.getElementById('printMethodToggle').getBoundingClientRect(); return [Math.round(r.left), Math.round(r.right), Math.round(r.height)];})()")
    pg.evaluate("document.getElementById('printMethodToggle').scrollIntoView({block:'center'})"); pg.wait_for_timeout(200)
    pg.tap('#printMethodToggle [data-pmt=no]'); pg.wait_for_timeout(200)
    hp = capture_print(pg, "printDish(encodeURIComponent(Object.keys(methodCrops)[0]))")
    pg.evaluate("document.getElementById('printMethodToggle').scrollIntoView({block:'center'})"); pg.wait_for_timeout(200)
    pg.screenshot(path=f'{SHOTS}/10-print-method-toggle-phone-390.png')
    check('phone 390px: toggle fits on screen, taps to No, dish print has no method images', box[0] >= 0 and box[1] <= 390 and box[2] >= 30 and not has_method(hp) and not errs, (box, errs))
    ctx8.close()
    # ---------- H. phone width ----------
    ctx, _ = new_ctx(b, 390, 844, is_mobile=True, has_touch=True, device_scale_factor=2); pg, errs = page_with_errs(ctx)
    pg.goto(URL); pg.wait_for_load_state('networkidle')
    up(pg, S + 'sample-recipes.pdf'); wait_status(pg, 'Sorted')
    sw = pg.evaluate('document.documentElement.scrollWidth')
    check('phone 390px: loads, sorts, no JS errors', not errs and pg.evaluate('ROWS.length') == 9, errs)
    check('phone 390px: version string visible in page', pg.get_attribute('meta[name="prepline-version"]', 'content') == ver)
    pg.screenshot(path=f'{SHOTS}/04-phone-390.png', full_page=False)
    enc = pg.evaluate("encName('Roast Pumpkin Salad')"); box = pg.locator(f'input[data-pax-loc="dish"][data-pax-dish="{enc}"]')
    box.tap(); pg.keyboard.press('Control+A'); pg.keyboard.press('Backspace'); pg.keyboard.type('60', delay=30); pg.wait_for_timeout(200)
    ph = ' | '.join(pg.locator('#resultsArea .section-table-wrap.no-print tr:has(td:has-text("Butternut pumpkin"))').all_inner_texts())
    check('phone 390px: typing pax (120 → 60) rescales live, 4 kg → 2 kg, keyboard stays in the box', '2 kg' in ph and pg.evaluate("document.activeElement.dataset.paxLoc") == 'dish', ph)
    pg.screenshot(path=f'{SHOTS}/07-phone-pax-live.png', full_page=False)
    results.append({'test': 'phone scrollWidth (info)', 'ok': True, 'info': sw})
    ctx.close(); b.close()

print('\nEXTERNAL REQUESTS:')
for h in sorted(ext):
    print(' ', h); [print('    ', u) for u in sorted(ext[h])]
bad = [h for h in ext if re.search(r'anthropic|openai|claude|generativelanguage|vision|ocr|tesseract|azure|aws|huggingface', h)]
if OFFLINE:
    check('offline: zero network requests (all network blocked, file:// page)', not net_attempts and not ext, (net_attempts[:5], sorted(ext)))
elif LIVE:
    allowed = {'cdnjs.cloudflare.com', 'fonts.googleapis.com', 'fonts.gstatic.com', urlparse(URL).hostname}
    check('live: outside requests only to the CDN, Google Fonts and the site itself (no Lemon Squeezy)', set(ext) <= allowed and not any('lemonsqueezy' in h for h in ext), sorted(ext))
check('network: no AI/OCR service hosts contacted', not bad, bad)
json.dump({'url': URL, 'version': ver, 'results': results, 'external': {h: sorted(v) for h, v in ext.items()}}, open(os.path.join(OUT, 'results.json'), 'w'), indent=1)
fails = [r for r in results if not r['ok']]
print(f'\n{len(results)-len(fails)}/{len(results)} passed')
sys.exit(1 if fails else 0)
